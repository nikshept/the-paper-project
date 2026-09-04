"""
Score -- reads marked answers off an already-dewarped photo.

LOGIC: for each detected option (a circle/checkbox region), count "ink"
pixels inside its box on the photo, and subtract the ink count at the
SAME box on a pixel-perfect blank render. A genuinely blank option
scores ~0 (whatever ink is there is in both images, so it cancels).
A real mark scores high (present in the photo, absent in the blank).
The highest-scoring option is the answer, UNLESS the numbers look
ambiguous (see CONFIG below) -- in which case it's flagged for a human
to check instead of guessed.

Everything a person would plausibly want to tune while debugging is a
named constant in the CONFIG block right below. Nothing else in this
file needs editing for normal tuning.
"""
from __future__ import annotations

import io

import cv2
import numpy as np
import pymupdf

from core.dewarp import SCALE


# ============================================================ CONFIG ===
# Every knob below is read only from this block -- change a number here,
# nothing else, restart the app, and re-score.

# --- Binarization (turns the photo/blank into pure black-ink-or-white,
#     BEFORE any box is even applied) -----------------------------------
BINARIZE_BLOCK_SIZE = 41   # Size (px) of the local neighborhood used to
                            # decide "ink or not" at each pixel. MUST be
                            # odd. Bigger = better at detecting a big
                            # solid mark fully (small blockSize can make
                            # the CENTER of a solid mark disappear, since
                            # its own local neighborhood is already dark).
                            # Too big = starts blending together lighting
                            # differences that are genuinely local (e.g.
                            # a shadow on one side of the page).
                            # Try: 25 (tighter) to 75 (looser).
BINARIZE_C = 10             # How much darker than local average a pixel
                            # must be to count as ink. Higher = stricter
                            # (fewer false positives from faint shadows,
                            # but a very light pencil mark may not clear
                            # the bar). Lower = more sensitive, more
                            # noise. Try: 5 (sensitive) to 30 (strict).

# --- Box position/size (where exactly ink gets counted) ----------------
BOX_MARGIN_PX = 0          # Shrinks or grows the sampled box from the
                            # detected circle's own edges, in pixels.
                            # POSITIVE = expand outward (risk: catches
                            # nearby table borders/dividers as false ink).
                            # NEGATIVE = shrink inward/inset (risk: over-
                            # weights whatever's at the box's own center,
                            # e.g. a printed digit inside the circle).
                            # 0 = exactly the detected box, no adjustment.
                            # This was expand(+4) -> inset(-15%) -> exact(0)
                            # across earlier tuning; exact was safest so far.
BOX_OFFSET_X_PX = 0        # Shifts EVERY sampled box sideways, in pixels.
BOX_OFFSET_Y_PX = 0        # Shifts EVERY sampled box up(-)/down(+), in
                            # pixels. Use this if the debug report/overlay
                            # shows boxes consistently off-center in one
                            # direction (a real, confirmed issue on some
                            # photos -- likely dewarp precision, not
                            # something margin/binarization can fix).
                            # Independent of BOX_MARGIN_PX -- offset moves
                            # the box, margin resizes it; use both together
                            # if needed.

# --- Decision thresholds (how the numbers above get turned into an
#     answer, or a flag for human review) -------------------------------
MIN_INK_THRESHOLD = 15     # An option's score must clear this to count
                            # as "marked at all". Below it on EVERY option
                            # -> no_mark_detected. Raise = stricter (more
                            # things read as blank). Lower = more
                            # sensitive to faint marks, but more noise-
                            # triggered false positives.
MIN_CONFIDENCE_GAP = 15    # The top-scoring option must beat the runner-
                            # up by at least this much, or it's flagged
                            # low_confidence. Raise = more flags. Lower =
                            # fewer flags, more risk of accepting a
                            # genuinely ambiguous mark.
MULTI_MARK_RATIO = 0.6     # An option only needs to reach this fraction
                            # of the TOP score to also count as "marked",
                            # triggering multiple_marks_detected if 2+
                            # options clear it. Raise toward 0.8-0.9 =
                            # much less trigger-happy on this flag (only
                            # near-equal marks count as "multiple").
                            # Lower = more sensitive to genuine double-
                            # marks, but more prone to flagging noise as
                            # a second mark.
# =========================================================================


def _binarize(gray):
    """Local (not global) threshold -- see BINARIZE_* above. Local
    matters because a real photo's lighting isn't even across the page;
    a single global cutoff would misread a shadowed blank area as ink."""
    return cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                  cv2.THRESH_BINARY_INV,
                                  blockSize=BINARIZE_BLOCK_SIZE, C=BINARIZE_C)


def build_blank_references(blank_pdf_bytes: bytes) -> dict:
    """Renders + binarizes the blank reference once, at the SAME pixel
    scale dewarped photos use, so a bbox in PDF points maps to the same
    pixel location in both -- no separate alignment step needed."""
    doc = pymupdf.open(stream=blank_pdf_bytes, filetype="pdf")
    refs = {}
    for pno, page in enumerate(doc):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(SCALE, SCALE))
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY) if pix.n >= 3 else img
        refs[pno] = _binarize(gray)
    doc.close()
    return refs


def binarize_dewarped(dewarped_bgr) -> "np.ndarray":
    gray = cv2.cvtColor(dewarped_bgr, cv2.COLOR_BGR2GRAY)
    return _binarize(gray)


def _bbox_to_px(bbox_pts: list) -> tuple:
    """PDF points -> pixels, THEN apply BOX_MARGIN_PX/BOX_OFFSET_*. This
    is the one place all three position/size knobs actually take effect."""
    x0, y0, x1, y1 = bbox_pts
    x0, y0, x1, y1 = x0 * SCALE, y0 * SCALE, x1 * SCALE, y1 * SCALE
    x0 = x0 - BOX_MARGIN_PX + BOX_OFFSET_X_PX
    x1 = x1 + BOX_MARGIN_PX + BOX_OFFSET_X_PX
    y0 = y0 - BOX_MARGIN_PX + BOX_OFFSET_Y_PX
    y1 = y1 + BOX_MARGIN_PX + BOX_OFFSET_Y_PX
    return int(x0), int(y0), int(x1), int(y1)


def _ink_score(bin_img, bbox_px: tuple) -> float:
    """Counts ink pixels in exactly the box _bbox_to_px produced."""
    x0, y0, x1, y1 = bbox_px
    h, w = bin_img.shape[:2]
    crop = bin_img[max(0, y0):min(h, y1), max(0, x0):min(w, x1)]
    if crop.size == 0:
        return 0.0
    return float(np.sum(crop > 0))


def score_question(dewarped_bin, blank_bin, option_elements: list) -> dict:
    """Scores one question. Returns answer_index (0-based, or None),
    confidence_gap, flagged, reason, and raw_diffs/raw_dewarped/
    raw_blank (per-option numbers, for debugging -- see the CSV report).

    Order of checks:
    1. Nothing clears MIN_INK_THRESHOLD -> no_mark_detected
    2. 2+ options clear MULTI_MARK_RATIO of the top -> multiple_marks_detected
    3. Top isn't ahead of runner-up by MIN_CONFIDENCE_GAP -> low_confidence
    4. Otherwise: confident single answer"""
    raw_dewarped, raw_blank, diffs = [], [], []
    bboxes_px = []
    for el in option_elements:
        bbox_px = _bbox_to_px(el["bbox"])
        bboxes_px.append(bbox_px)
        dw = _ink_score(dewarped_bin, bbox_px)
        bl = _ink_score(blank_bin, bbox_px)
        raw_dewarped.append(dw)
        raw_blank.append(bl)
        diffs.append(dw - bl)

    if not diffs:
        return {"answer_index": None, "confidence_gap": 0, "flagged": True,
                "reason": "no_options_in_template", "raw_diffs": diffs,
                "raw_dewarped": raw_dewarped, "raw_blank": raw_blank, "bboxes_px": bboxes_px}

    ranked = sorted(diffs, reverse=True)
    top = ranked[0]
    gap = (ranked[0] - ranked[1]) if len(ranked) > 1 else top
    n_marked = sum(1 for d in diffs if d >= MIN_INK_THRESHOLD and d >= MULTI_MARK_RATIO * top)

    result = {"raw_diffs": diffs, "raw_dewarped": raw_dewarped, "raw_blank": raw_blank,
              "bboxes_px": bboxes_px, "confidence_gap": gap}

    if top < MIN_INK_THRESHOLD:
        result.update(answer_index=None, flagged=True, reason="no_mark_detected")
    elif n_marked >= 2:
        result.update(answer_index=int(np.argmax(diffs)), flagged=True, reason="multiple_marks_detected")
    else:
        answer_index = int(np.argmax(diffs))
        flagged = gap < MIN_CONFIDENCE_GAP
        result.update(answer_index=answer_index, flagged=flagged,
                      reason="low_confidence" if flagged else None)
    return result


def score_response(dewarped_images: dict, blank_refs: dict, confirmed_template: list) -> dict:
    """One respondent. Returns {question_index: score_dict_or_None} --
    None for a page not available for this respondent, or an open-text
    question (needs manual transcription, not ink-diff)."""
    results = {}
    for i, g in enumerate(confirmed_template):
        if g["page"] not in dewarped_images:
            results[i] = None
            continue
        option_elements = g["option_elements"]
        # Falls back to the old implicit heuristic if question_type is missing
        # (a template saved before that field existed).
        q_type = g.get("question_type") or ("open_text" if len(option_elements) == 1 else "mcq")
        if q_type == "open_text":
            results[i] = None
            continue
        dewarped_bin = binarize_dewarped(dewarped_images[g["page"]])
        blank_bin = blank_refs.get(g["page"])
        if blank_bin is None:
            results[i] = None
            continue
        results[i] = score_question(dewarped_bin, blank_bin, option_elements)
    return results


# ============================================================ EXPORTS ==

def build_export_xlsx(scored_responses: dict, confirmed_template: list, corrections: dict) -> bytes:
    """The researcher-facing spreadsheet. One row per respondent, one
    column per question. Yellow = flagged/unreviewed, green = corrected
    in-app. Each cell's comment shows why it was flagged, for a quick
    check without needing the full debug report below."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from openpyxl.comments import Comment

    FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    REVIEWED_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Responses"

    def col_header(i, g):
        return g.get("label") or (g.get("question_text_guess") or f"Q{i + 1}")[:40]

    headers = ["tracking_code"] + [col_header(i, g) for i, g in enumerate(confirmed_template)]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for code in sorted(scored_responses.keys()):
        row_idx = ws.max_row + 1
        ws.cell(row=row_idx, column=1, value=code)
        for qi in range(len(confirmed_template)):
            score = scored_responses[code].get(qi)
            correction = corrections.get((code, qi))
            cell = ws.cell(row=row_idx, column=qi + 2)

            comment_lines = []
            if score is not None:
                diffs_str = ", ".join(f"{d:.0f}" for d in score.get("raw_diffs", []))
                comment_lines.append(f"Detected reason: {score.get('reason') or 'confident'}")
                comment_lines.append(f"Confidence gap: {score.get('confidence_gap', 0):.0f}")
                comment_lines.append(f"Raw ink-diff scores: [{diffs_str}]")
            elif correction is None:
                comment_lines.append("No photo available for this page, or open-text "
                                     "(needs manual transcription).")

            if correction is not None:
                cell.value = correction
                cell.fill = REVIEWED_FILL
                if comment_lines:
                    comment_lines.insert(0, "Corrected in review (was: auto-detected).")
            elif score is None:
                cell.value = ""
                cell.fill = FLAG_FILL
            else:
                if score["answer_index"] is None:
                    cell.value = "NA"
                elif score["reason"] == "multiple_marks_detected":
                    cell.value = "MULT"
                else:
                    cell.value = score["answer_index"] + 1
                if score["flagged"]:
                    cell.fill = FLAG_FILL

            if comment_lines:
                cell.comment = Comment("\n".join(comment_lines), "Paper Survey Tool")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_debug_report_xlsx(scored_responses: dict, confirmed_template: list) -> bytes:
    """One row per (respondent, option) -- every individual box scored,
    not just the winning answer per question. Includes the exact pixel
    box used, its raw dewarped/blank/diff ink counts, and the CONFIG
    values active when it was generated (so a downloaded report is
    self-describing even after you've since changed the numbers).
    Downloadable any time after scoring, independent of review progress."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    TOP_PICK_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Debug"

    ws.append(["config: BLOCK_SIZE", BINARIZE_BLOCK_SIZE, "C", BINARIZE_C,
              "BOX_MARGIN_PX", BOX_MARGIN_PX, "BOX_OFFSET_X_PX", BOX_OFFSET_X_PX,
              "BOX_OFFSET_Y_PX", BOX_OFFSET_Y_PX, "MIN_INK_THRESHOLD", MIN_INK_THRESHOLD,
              "MIN_CONFIDENCE_GAP", MIN_CONFIDENCE_GAP, "MULTI_MARK_RATIO", MULTI_MARK_RATIO])
    ws.append([])
    ws.append(["tracking_code", "question_index", "label", "page", "option_index",
              "box_x0_px", "box_y0_px", "box_x1_px", "box_y1_px",
              "ink_dewarped", "ink_blank", "ink_diff",
              "is_top_pick", "question_reason", "question_confidence_gap"])
    for cell in ws[3]:
        cell.font = Font(bold=True)

    for code in sorted(scored_responses.keys()):
        for qi, g in enumerate(confirmed_template):
            score = scored_responses[code].get(qi)
            if score is None or "bboxes_px" not in score:
                continue  # open-text or missing page -- no boxes to report
            top_idx = score.get("answer_index")
            for oi, bbox_px in enumerate(score["bboxes_px"]):
                is_top = oi == top_idx
                ws.append([
                    code, qi, g.get("label", ""), g.get("page", ""), oi,
                    bbox_px[0], bbox_px[1], bbox_px[2], bbox_px[3],
                    f"{score['raw_dewarped'][oi]:.0f}", f"{score['raw_blank'][oi]:.0f}",
                    f"{score['raw_diffs'][oi]:.0f}",
                    is_top, score.get("reason") or "confident",
                    f"{score.get('confidence_gap', 0):.0f}",
                ])
                if is_top:
                    for cell in ws[ws.max_row]:
                        cell.fill = TOP_PICK_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
