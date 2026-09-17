import cv2, json, copy
import pymupdf 
import numpy as np
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
import getResponses
import streamlit as st

### Fetch selected pages of a pdf document as images
def fetch_images (document, selected_pages, page_dpi):
    images = {}

    for i, page in enumerate(document):
        if i not in selected_pages:
            continue
        pix = page.get_pixmap(dpi=page_dpi)
        arr = np.frombuffer(pix.samples, dtype="uint8").reshape(pix.height, pix.width, pix.n)
        images[f"page{i+1}"] = cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)

    return images

### Crop images by box dimensions
def img_cropper (image, box):
    h, w = image.shape[:2]
    x1, y1, x2, y2 = box["x1"], box["y1"], box["x2"], box["y2"]

    croppedImg = image[int(y1*h):int(y2*h), int(x1*w):int(x2*w)]

    return croppedImg

### Get QR Code function
def get_QR (croppedboxImage):
    qrID, _, _ = cv2.QRCodeDetector().detectAndDecode(croppedboxImage)
    if qrID == "":
        res, points = cv2.wechat_qrcode.WeChatQRCode().detectAndDecode(croppedboxImage)
        qrID = res[0] if res else ""

    return qrID

### Process the pages and get results (responses and scored images)
def process_pages(template_file, pdf_file, selected_pages, page_dpi, output_json_path=None, save_to_disk=False, debug=False):
    # input template and checks whether there is a QR box and some response boxes
    template = json.load(template_file)

    if not any(box["box_type"] == "QR_Box" for box in template):
        print("No QR box in template")
        st.error("No QR box found in template. Please check your template.")
        return []
    if not any(box["box_type"] == "Response_Box" for box in template):
        print("No Response box in template")
        st.error("No Response box found in template. Please check your template.")
        return []

    # convert pdf into images
    doc = pymupdf.open(stream=pdf_file.read(), filetype="pdf")
    images = fetch_images(doc, selected_pages, page_dpi)

    all_results = []
    st.session_state.all_results_images = {}
    st.session_state.all_cropped_boxes = {}

    # cycle through page images and get answers
    for key, img in images.items():
        boxes = copy.deepcopy(template)
        page_qr_id = ""
        print(f"Investigating {key}")

        for box in boxes:
            if box["box_type"] == "QR_Box":
                croppedQRimg = img_cropper(img, box)
                page_qr_id = get_QR(croppedQRimg)
                if page_qr_id != "":
                    print(f"QR Box found in {key} and identified as {page_qr_id}")

        if page_qr_id == "":
            print(f"Skipping {key}: QR not identified")
            continue

        for i, box in enumerate(boxes):
            if box["box_type"] == "Response_Box":
                croppedResimg = img_cropper(img, box)
                tag = f"{key}_{page_qr_id}_Box{i}"
                box["tag"] = tag
                st.session_state.all_cropped_boxes[tag] = croppedResimg
                if debug:
                    cv2.imwrite(f"output/00_cropped_box_{tag}.png", croppedResimg)

                results, imgResults = getResponses.get_Responses(croppedResimg, box["bperRow"], tag, debug)

                for label, (value, gap_ratio, flag, source, flagReason) in zip(box["questionlabels"].keys(), results):
                    box["questionlabels"][label] = {"value": value, "confidence_ratio": gap_ratio, "flagged": flag, "source": source, "flagReason": flagReason}

                st.session_state.all_results_images[tag] = imgResults

                valsCount = len(results)
                flagCount = sum(1 for (_, _, flag, _, _) in results if flag)
                print(f"Box{i}: Received {valsCount} values and {flagCount} were flagged.")

        all_results.append({"page": key, "qr_ID": page_qr_id, "boxes": boxes})

    if save_to_disk:
        json.dump(all_results, open(output_json_path, "w"), indent=2)
    return all_results

### Build excel sheet based on results
def build_excel(all_results, output_path):
    RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    wb = openpyxl.Workbook()
    ws = wb.active

    all_labels = []
    for page in all_results:
        for box in page["boxes"]:
            for label in box.get("questionlabels", {}):
                if label not in all_labels:
                    all_labels.append(label)

    ws.append(["page", "qr_ID"] + all_labels)

    for page in all_results:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=page["page"])
        ws.cell(row=r, column=2, value=page["qr_ID"])
        for box in page["boxes"]:
            if box["box_type"] != "Response_Box":
                continue
            for label, ans in box.get("questionlabels", {}).items():
                if not isinstance(ans, dict):
                    continue
                col = 3 + all_labels.index(label)
                cell = ws.cell(row=r, column=col, value=ans.get("value"))
                cell.comment = Comment(f"confidence_ratio: {ans.get('confidence_ratio'):.0%}\nflagged: {ans.get('flagged')}\nsource: {ans.get('source')}", "cropper")
                if ans.get("flagged"):
                    cell.fill = RED_FILL

    wb.save(output_path)

if __name__ == "__main__":
    all_results = process_pages(
        template_file=open("input/page1_template.json", "rb"),
        pdf_file=open("input/merged_input.pdf", "rb"),
        selected_pages=[0],
        page_dpi=280,
        output_json_path="output/all_pages_results.json",
        save_to_disk=True,
    )
    build_excel(all_results, "output/responses.xlsx")   