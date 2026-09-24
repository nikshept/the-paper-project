import cv2, json, copy, pymupdf, io
import numpy as np
import openpyxl
from dewarp import dewarp_image_A4
import getResponses
import streamlit as st
import streamlit.logger
import os
streamlit.logger.set_log_level("ERROR")
os.makedirs("output", exist_ok=True)

################################ HELPERS #############################

### Renders pages of a pdf as a set of image files
def pdf_to_image_files(pdf_file, dpi=280):
    doc = pymupdf.open(stream=pdf_file.read(), filetype="pdf")
    image_files = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi)
        fake_file = io.BytesIO(pix.tobytes("png"))
        fake_file.name = f"page{i+1}.png"
        image_files.append(fake_file)
    return image_files

### Crop images by box dimensions
def img_cropper (image, box):
    h, w = image.shape[:2]
    x1, y1, x2, y2 = box["x1"], box["y1"], box["x2"], box["y2"]

    croppedImg = image[int(y1*h):int(y2*h), int(x1*w):int(x2*w)]

    return croppedImg

### Get QR Code function
def get_QR (croppedboxImage):
    qrID, _, _ = cv2.QRCodeDetector().detectAndDecode(croppedboxImage)
    if qrID == "":
        res, points = cv2.wechat_qrcode.WeChatQRCode().detectAndDecode(croppedboxImage)
        qrID = res[0] if res else ""

    return qrID

################################ MAIN PIPELINE #############################

### Validate if the template has QR boxes and response boxes
def validate_template(template, debug=False):
    if not any(box["box_type"] == "QR_Box" for box in template):
        msg = "No QR box found in template. Please re-upload a valid template."
        print(msg) if debug else st.error(msg)
        return False
    if not any(box["box_type"] in ("Response_Box", "Text_Box") for box in template):
        msg = "No Response or Text box found in template. Please re-upload a valid template."
        print(msg) if debug else st.error(msg)
        return False
    return True

### Validate the images for successful dewarp, and reads QRs
def validate_images(image_files, template, debug=False):
    qr_box = next(box for box in template if box["box_type"] == "QR_Box")
    flagged_images = []
    dewarped_images = {}
    qr_ids = {}

    for img_file in image_files:
        file_bytes = np.frombuffer(img_file.read(), dtype=np.uint8)
        original = cv2.imdecode(file_bytes, cv2.IMREAD_COLOR)
        img_file.seek(0)  # reset -- dewarp_image_A4 reads this same stream again below

        if original.shape[1] < 1500:
            flagged_images.append(f"{img_file.name} (low quality)")
            continue

        dewarped = dewarp_image_A4(img_file, debug=debug)
        if dewarped is None:
            flagged_images.append(f"{img_file.name} (dewarp failed)")
            continue

        cropped_qr = img_cropper(dewarped, qr_box)
        qr_id = get_QR(cropped_qr)
        if qr_id == "":
            flagged_images.append(f"{img_file.name} (QR not detected)")
            continue

        dewarped_images[os.path.basename(img_file.name)] = dewarped
        qr_ids[os.path.basename(img_file.name)] = qr_id

    if flagged_images:
        msg = f"Please re-upload after improving the quality of: {flagged_images}"
        print(msg) if debug else st.error(msg)

    return flagged_images, dewarped_images, qr_ids

### Process the pages and get results (responses and scored images)
def extract_results(template, dewarped_images, qr_ids, debug=False):
    all_results = []
    st.session_state.all_results_images = {}
    st.session_state.text_results_images = {}
    st.session_state.all_cropped_boxes = {}

    # cycle through page images and get answers
    for key, img in dewarped_images.items():
        boxes = copy.deepcopy(template)
        page_qr_id = qr_ids[key]
        print(f"Investigating {key}, QR: {page_qr_id}")

        for i, box in enumerate(boxes):
            if box["box_type"] == "Response_Box":
                croppedResimg = img_cropper(img, box)
                tag = f"{key}_{page_qr_id}_Response_Box{i}"
                box["tag"] = tag
                st.session_state.all_cropped_boxes[tag] = croppedResimg
                if debug:
                    cv2.imwrite(f"output/00_cropped_{tag}.png", croppedResimg)

                results, imgResults = getResponses.get_Responses(croppedResimg, box["bperRow"], tag, debug)

                for label, (value, gap_ratio, flag, source, flagReason) in zip(box["questionlabels"].keys(), results):
                    box["questionlabels"][label] = {"model_value": value, "confidence_ratio": gap_ratio, "flagged": flag, "source": source, "flagReason": flagReason}

                st.session_state.all_results_images[tag] = imgResults

                valsCount = len(results)
                flagCount = sum(1 for (_, _, flag, _, _) in results if flag)
                print(f"Box{i}: Received {valsCount} values and {flagCount} were flagged.")

            elif box["box_type"] == "Text_Box":
                croppedTextimg = img_cropper(img, box)
                tag = f"{key}_{page_qr_id}_Text_Box{i}"
                box["tag"] = tag
                st.session_state.all_cropped_boxes[tag] = croppedTextimg
                # call a function to process the text in the text image
                st.session_state.text_results_images[tag] = croppedTextimg

                if debug:
                    cv2.imwrite(f"output/00_cropped_{tag}.png", croppedTextimg)            

        all_results.append({"page": key, "qr_ID": page_qr_id, "boxes": boxes})

    

    return all_results

### Build excel sheet of responses
def build_responses_excel(all_results, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active

    all_labels = []
    for page in all_results:
        for box in page["boxes"]:
            for label in box.get("questionlabels", {}):
                if label not in all_labels:
                    all_labels.append(label)

    ws.append(["page", "qr_ID"] + all_labels)

    for page in all_results:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=page["page"])
        ws.cell(row=r, column=2, value=page["qr_ID"])
        for box in page["boxes"]:
            if box["box_type"] == "QR_Box":
                continue
            for label, ans in box.get("questionlabels", {}).items():
                if not isinstance(ans, dict):
                    continue
                col = 3 + all_labels.index(label)
                ws.cell(row=r, column=col, value=ans.get("user_value"))

    wb.save(output_path)

### Build excel sheet for model evaluation
def build_eval_excel(all_results, output_path):
    wb = openpyxl.Workbook()

    # ---- sheet 1: one row per question, raw data ----
    ws = wb.active
    ws.title = "Answers"
    ws.append(["box_tag", "question_label", "model_value", "user_value", "source", "flagReason", "confidence_ratio"])

    # correct/wrong counts per flagReason, built up as we scan every answer
    tally = {}  # {flagReason: [correct_count, wrong_count]}

    for page in all_results:
        for box in page["boxes"]:

            # skip QR boxes
            if box["box_type"] == "QR_Box":
                continue

            for label, ans in box.get("questionlabels", {}).items():
                # skip untouched template placeholders
                if not isinstance(ans, dict):
                    continue

                model_value = ans.get("model_value")
                user_value = ans.get("user_value")
                flagReason = ans.get("flagReason")

                ws.append([
                    box.get("tag"),
                    label,
                    model_value,
                    user_value,
                    ans.get("source"),
                    flagReason,
                    ans.get("confidence_ratio"),
                ])

                # Text_Box has no model_value at the moment, skip their eval
                if flagReason is None:
                    continue

                # Tally response bubbles' correct and wrongly predicted vals
                tally.setdefault(flagReason, [0, 0])

                if flagReason == "low confidence": # is correct if the real answer is either NA / MULT, if it was a num, its wrong
                    is_correct = user_value in ("NA", "MULT")
                elif flagReason == "high confidence" or flagReason == "medium confidence":
                    is_correct = str(model_value) == str(user_value) # is correct if the prediction matches user's confirmation

                if is_correct:
                    tally[flagReason][0] += 1  # correct answers count +1
                else:
                    tally[flagReason][1] += 1  # wrong answers count +1

    # ---- sheet 2: accuracy/coverage summary ----
    stats = wb.create_sheet("Stats")
    stats.append(["Confidence Level", "Correct", "Wrong", "Total", "Accuracy"])

    committed = [0, 0]  # [correct, wrong] -- "high"/"medium" confidence = model committed to an answer
    overall = [0, 0]    # same, but across every flagReason including "low confidence"

    for reason in ["high confidence", "medium confidence", "low confidence"]:
        if reason not in tally:
            continue

        # compute per reason accuracy = fraction of correct answers
        correct, wrong = tally[reason]
        total = correct + wrong
        accuracy = correct / total if total else 0
        stats.append([reason, correct, wrong, total, accuracy])

        # compute overall accuracy and coverage = fraction of committed/predicted answers
        overall[0] += correct
        overall[1] += wrong
        if reason != "low confidence":
            committed[0] += correct
            committed[1] += wrong

    overall_total = sum(overall)
    committed_total = sum(committed)

    stats.append(["Overall", overall[0], overall[1], overall_total, overall[0] / overall_total if overall_total else 0])
    stats.append([])
    stats.append(["Coverage (high/medium confidence predictions)", "", "", "", committed_total / overall_total if overall_total else 0])
    stats.append(["Selective accuracy (of high & medium conf predictions)", "", "", "", committed[0] / committed_total if committed_total else 0])

    wb.save(output_path)

### Top level function: First validates, and then extracts results
def process_images(image_files, template_file, debug=False):
    template = json.load(template_file)

    if not validate_template(template, debug=debug):
        return None

    flagged_images, dewarped_images, qr_ids = validate_images(image_files, template, debug=debug)
    if flagged_images:
        return None

    return extract_results(template, dewarped_images, qr_ids, debug=debug)

if __name__ == "__main__":
    image_files = [open(f"input/{name}", "rb") for name in ["image17.jpg"]]
    template_file = open("input/page4_template.json", "rb")

    all_results = process_images(image_files, template_file, debug=True)
    if all_results:
        build_responses_excel(all_results, "output/responses.xlsx")
        json.dump(all_results, open("output/all_pages_results.json", "w"), indent=2)